import os
from datetime import datetime
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from models import CompetitorSnapshot, PriceDelta, DiffReport

class ExcelManager:
    """
    Manages the competitor_price_tracker.xlsx workbook with executive financial formatting,
    KPI summary blocks, auto-fit columns, currency formatting, and color-coded audit logs.
    """
    def __init__(self, excel_path: str = "./data/competitor_price_tracker.xlsx"):
        self.excel_path = excel_path
        os.makedirs(os.path.dirname(os.path.abspath(self.excel_path)), exist_ok=True)
        self._ensure_workbook_exists()

    def _ensure_workbook_exists(self):
        if not os.path.exists(self.excel_path):
            wb = openpyxl.Workbook()
            ws_overview = wb.active
            ws_overview.title = "Current Pricing Overview"
            ws_history = wb.create_sheet(title="Price Change History")
            try:
                wb.save(self.excel_path)
            except Exception:
                pass

    def _apply_thin_borders(self, ws, min_row, max_row, min_col, max_col):
        thin_side = Side(border_style="thin", color="E2E8F0")
        border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                ws.cell(row=r, column=c).border = border

    def update_with_diff_report(self, snapshot: CompetitorSnapshot, diff_report: DiffReport):
        try:
            wb = openpyxl.load_workbook(self.excel_path)
        except Exception:
            wb = openpyxl.Workbook()

        now_str = datetime.now().strftime("%d %b %Y - %I:%M %p")

        # ----------------------------------------------------
        # 1. FORMAT SHEET 1: CURRENT PRICING OVERVIEW
        # ----------------------------------------------------
        if "Current Pricing Overview" not in wb.sheetnames:
            ws_overview = wb.create_sheet(title="Current Pricing Overview", index=0)
        else:
            ws_overview = wb["Current Pricing Overview"]

        # Read existing data rows (if any) to preserve other competitors
        existing_rows = []
        if ws_overview.max_row >= 8:
            for r in ws_overview.iter_rows(min_row=8, values_only=True):
                if r and r[0] and r[0] != snapshot.competitor_name and not str(r[0]).startswith("MARGIN"):
                    existing_rows.append(r)

        # Clear and rebuild overview sheet cleanly
        ws_overview.delete_rows(1, ws_overview.max_row + 10)
        ws_overview.views.sheetView[0].showGridLines = True

        # Sheet Title Banner
        ws_overview["A1"] = "MARGIN GUARD | Competitive Intelligence Overview"
        ws_overview["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="0F172A")
        
        ws_overview["A2"] = f"Apex Wellness & Aesthetics  |  Last Synced: {now_str}  |  Source: Monitored Competitor Endpoints"
        ws_overview["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

        # Summary KPI Mini-Boxes (Row 4 to 5)
        kpis = [
            ("COMPETITOR MONITORED", snapshot.competitor_name, "0F172A", "F8FAFC"),
            ("SERVICES TRACKED", f"{len(snapshot.items)} Items", "1E3A8A", "EFF6FF"),
            ("PRICE HIKES DETECTED", f"{len(diff_report.price_increases)} Active", "991B1B", "FEF2F2"),
            ("PRICE DROPS DETECTED", f"{len(diff_report.price_decreases)} Promos", "166534", "F0FDF4")
        ]
        
        col_start = 1
        for title, val, text_col, bg_col in kpis:
            c_title = ws_overview.cell(row=4, column=col_start, value=title)
            c_title.font = Font(name="Segoe UI", size=8, bold=True, color="64748B")
            c_title.fill = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")
            c_title.alignment = Alignment(horizontal="left", vertical="center")
            
            c_val = ws_overview.cell(row=5, column=col_start, value=val)
            c_val.font = Font(name="Segoe UI", size=11, bold=True, color=text_col)
            c_val.fill = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")
            c_val.alignment = Alignment(horizontal="left", vertical="center")
            
            thin_s = Side(border_style="thin", color="CBD5E1")
            c_title.border = Border(left=thin_s, right=thin_s, top=thin_s)
            c_val.border = Border(left=thin_s, right=thin_s, bottom=thin_s)
            col_start += 1

        # Table Headers (Row 7)
        overview_headers = ["Competitor Name", "Category", "Service / Treatment", "Current Price", "Unit / Scope", "Verification Status", "Last Verified"]
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
        
        for idx, h_text in enumerate(overview_headers, start=1):
            c = ws_overview.cell(row=7, column=idx, value=h_text)
            c.fill = header_fill
            c.font = header_font
            align_h = "right" if "Price" in h_text else ("center" if "Last" in h_text or "Status" in h_text else "left")
            c.alignment = Alignment(horizontal=align_h, vertical="center")
        ws_overview.row_dimensions[7].height = 24

        # Populate Overview Rows
        all_overview_rows = existing_rows + [
            (snapshot.competitor_name, it.category, it.name, it.price, it.unit or "session", "Verified", now_str)
            for it in snapshot.items
        ]

        for r_idx, row_data in enumerate(all_overview_rows, start=8):
            zebra_bg = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
            row_fill = PatternFill(start_color=zebra_bg, end_color=zebra_bg, fill_type="solid")
            
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws_overview.cell(row=r_idx, column=c_idx, value=val)
                cell.font = Font(name="Segoe UI", size=9.5, color="1E293B")
                cell.fill = row_fill
                
                if c_idx == 4:
                    cell.number_format = '$#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
                elif c_idx in [6, 7]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            ws_overview.row_dimensions[r_idx].height = 20

        self._apply_thin_borders(ws_overview, 7, ws_overview.max_row, 1, len(overview_headers))
        ws_overview.freeze_panes = "A8"

        # ----------------------------------------------------
        # 2. FORMAT SHEET 2: PRICE CHANGE HISTORY
        # ----------------------------------------------------
        if "Price Change History" not in wb.sheetnames:
            ws_history = wb.create_sheet(title="Price Change History")
        else:
            ws_history = wb["Price Change History"]

        ws_history.views.sheetView[0].showGridLines = True

        # Check if history has clean banner headers
        if ws_history.cell(row=1, column=1).value != "MARGIN GUARD | Historical Price Delta & Audit Log":
            ws_history.delete_rows(1, ws_history.max_row + 10)
            ws_history["A1"] = "MARGIN GUARD | Historical Price Delta & Audit Log"
            ws_history["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="0F172A")
            ws_history["A2"] = "Complete chronological audit log of all detected competitor pricing shifts, promotions, and menu adjustments."
            ws_history["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

            history_headers = ["Timestamp", "Competitor Name", "Category", "Service / Item", "Previous Price", "New Price", "Net Delta ($)", "Change (%)", "Movement Type"]
            hist_header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
            
            for idx, h_text in enumerate(history_headers, start=1):
                c = ws_history.cell(row=5, column=idx, value=h_text)
                c.fill = hist_header_fill
                c.font = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
                align_h = "right" if "Price" in h_text or "Delta" in h_text or "%" in h_text else ("center" if "Timestamp" in h_text or "Type" in h_text else "left")
                c.alignment = Alignment(horizontal=align_h, vertical="center")
            ws_history.row_dimensions[5].height = 24
            ws_history.freeze_panes = "A6"

        # Append new changes to history
        red_bg = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid") # Price hike
        green_bg = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid") # Price drop
        blue_bg = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid") # New item
        
        red_font = Font(name="Segoe UI", size=9.5, bold=True, color="991B1B")
        green_font = Font(name="Segoe UI", size=9.5, bold=True, color="166534")
        blue_font = Font(name="Segoe UI", size=9.5, bold=True, color="1D4ED8")
        neutral_font = Font(name="Segoe UI", size=9.5, color="1E293B")

        changes = diff_report.price_increases + diff_report.price_decreases + diff_report.new_items + diff_report.discontinued_items
        for delta in changes:
            curr_r = ws_history.max_row + 1
            if curr_r < 6:
                curr_r = 6
                
            ts_str = delta.detected_at.strftime("%Y-%m-%d %H:%M")
            old_p = delta.old_price if delta.old_price is not None else ""
            new_p = delta.new_price if delta.new_price is not None else ""
            diff_amt = delta.delta_amount if delta.delta_amount is not None else ""
            diff_pct = (delta.delta_percentage / 100.0) if delta.delta_percentage is not None else ""
            m_type = delta.change_type.replace("_", " ")

            row_data = [ts_str, delta.competitor_name, delta.category, delta.item_name, old_p, new_p, diff_amt, diff_pct, m_type]
            
            if delta.change_type == "PRICE_INCREASE":
                r_fill = red_bg
                r_font = red_font
            elif delta.change_type == "PRICE_DECREASE":
                r_fill = green_bg
                r_font = green_font
            elif delta.change_type == "NEW_ITEM":
                r_fill = blue_bg
                r_font = blue_font
            else:
                r_fill = PatternFill(start_color="F8FAFC", fill_type="solid")
                r_font = neutral_font

            for c_idx, val in enumerate(row_data, start=1):
                cell = ws_history.cell(row=curr_r, column=c_idx, value=val)
                cell.fill = r_fill
                cell.font = r_font
                
                if c_idx in [5, 6]:
                    if isinstance(val, (int, float)):
                        cell.number_format = '$#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx == 7:
                    if isinstance(val, (int, float)):
                        cell.number_format = '+$#,##0.00;-$#,##0.00;"$0.00"'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx == 8:
                    if isinstance(val, (int, float)):
                        cell.number_format = '+0.0%;-0.0%;"0.0%"'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx in [1, 9]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            ws_history.row_dimensions[curr_r].height = 20

        if ws_history.max_row >= 5:
            self._apply_thin_borders(ws_history, 5, ws_history.max_row, 1, 9)

        # ----------------------------------------------------
        # 3. AUTO-FIT COLUMN WIDTHS ACROSS ALL SHEETS
        # ----------------------------------------------------
        for ws in [ws_overview, ws_history]:
            for col in ws.columns:
                max_len = 0
                for cell in col:
                    if cell.row in [1, 2]: # Ignore long title banner rows
                        continue
                    if cell.value:
                        val_str = str(cell.value)
                        max_len = max(max_len, len(val_str))
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 5, 16)

        try:
            wb.save(self.excel_path)
        except PermissionError:
            fallback_path = self.excel_path.replace(".xlsx", f"_{datetime.now().strftime('%H%M%S')}.xlsx")
            wb.save(fallback_path)
            print(f"      [NOTE] Main Excel file was open/locked; saved update to: {fallback_path}")
            return fallback_path

        return self.excel_path
